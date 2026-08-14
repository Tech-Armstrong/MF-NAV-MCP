"""
writer.py
---------
Writes the enriched holdings and validation results to an xlsx workbook.

Sheets produced
---------------
1. Holdings Enriched  — all equity rows with AMFI data, colour-coded by match method
2. Alias Report       — one row per unique stock, showing how it was matched
3. Weight Validation  — per-fund % sum check with pass/fail status
4. Legend             — colour key
5. Fund Profile       — one row per fund: market-cap split + sector breakdown
"""

from pathlib import Path
from typing import List, Dict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ── Colour palette ─────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")

FILL = {
    "exact-norm":   PatternFill("solid", fgColor="E2EFDA"),   # green
    "manual-alias": PatternFill("solid", fgColor="FFF2CC"),   # yellow
    "fuzzy-high":   PatternFill("solid", fgColor="DDEBF7"),   # blue
    "no-match":     PatternFill("solid", fgColor="FCE4D6"),   # red
    "manual-none":  PatternFill("solid", fgColor="F2F2F2"),   # grey
}

STATUS_FILL = {
    "OK":         PatternFill("solid", fgColor="E2EFDA"),
    "INCOMPLETE": PatternFill("solid", fgColor="FCE4D6"),
    "OVER 100":   PatternFill("solid", fgColor="FFF2CC"),
}

# ── Column order for main sheet ────────────────────────────────────────────
MAIN_COLS = [
    "Cap Category", "Fund", "Stock", "Holding %", "As of",
    "AMFI Name", "ISIN", "Match Method", "Match Score",
    "Mkt Cap Cat", "Mkt Cap ₹Cr", "Industry", "Sector",
]

ALIAS_COLS  = ["Stock", "AMFI Name", "ISIN", "Match Method", "Match Score",
               "Mkt Cap Cat", "Mkt Cap ₹Cr", "Industry", "Sector"]

VALID_COLS  = ["Fund", "Equity %", "Non-Equity %", "Sum of Holding %",
               "Deviation", "Status", "Tolerance (±%)"]

# ── Fund Profile ───────────────────────────────────────────────────────────
# One row per fund: market-cap split, then one column per sector. Both blocks
# sum to the fund's total equity weight (not to 100 — the rest is cash, repo
# and derivatives, which are not equity holdings).
CAP_ORDER    = ["Large Cap", "Mid Cap", "Small Cap"]
CAP_COLS     = [f"% {c}" for c in CAP_ORDER] + ["% Unclassified", "Cap Total"]
UNCLASSIFIED = "Unclassified Sector"


def build_fund_profile(enriched: List[Dict]) -> tuple[List[Dict], List[str]]:
    """
    Aggregate enriched rows into one row per fund.

    Returns (rows, sector_columns). Sector columns are derived from the data
    so a new AMFI sector appears automatically without a code change.
    """
    sectors = sorted({(r.get("Sector") or "").strip()
                      for r in enriched if (r.get("Sector") or "").strip()})
    sector_cols = sectors + [UNCLASSIFIED]

    funds: Dict[str, Dict] = {}
    for r in enriched:
        fund = r.get("Fund") or ""
        if not fund:
            continue
        prof = funds.setdefault(fund, {"Fund": fund})

        pct = r.get("Holding %") or 0.0

        # market-cap bucket
        cap = (r.get("Mkt Cap Cat") or "").strip()
        cap_key = f"% {cap}" if cap in CAP_ORDER else "% Unclassified"
        prof[cap_key] = prof.get(cap_key, 0.0) + pct

        # sector bucket
        sec = (r.get("Sector") or "").strip() or UNCLASSIFIED
        prof[sec] = prof.get(sec, 0.0) + pct

    rows = []
    for fund in sorted(funds):
        p = funds[fund]
        p["Cap Total"]    = round(sum(p.get(c, 0.0) for c in CAP_COLS[:-1]), 5)
        p["Sector Total"] = round(sum(p.get(s, 0.0) for s in sector_cols), 5)
        for k, v in p.items():
            if isinstance(v, float):
                p[k] = round(v, 5)
        rows.append(p)

    return rows, sector_cols


def _write_sheet(ws, cols: List[str], rows: List[Dict], row_fill_key: str = None):
    """Generic sheet writer: header row then data rows."""
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([row.get(c) for c in cols])
        if row_fill_key and row.get(row_fill_key) in FILL:
            fill = FILL[row[row_fill_key]]
            for cell in ws[ws.max_row]:
                cell.fill = fill

    _autowidth(ws)


def _autowidth(ws, max_width: int = 45):
    for col in ws.columns:
        w = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(w + 2, max_width)


def write_output(
    enriched:        List[Dict],
    alias_rows:      List[Dict],
    validation_rows: List[Dict],
    output_path:     Path,
):
    """
    Write all output sheets to a single xlsx workbook.
    """
    wb = openpyxl.Workbook()

    # ── Sheet 1: Holdings Enriched ─────────────────────────────────────────
    ws_main = wb.active
    ws_main.title = "Holdings Enriched"
    _write_sheet(ws_main, MAIN_COLS, enriched, row_fill_key="Match Method")

    # format Mkt Cap ₹Cr as number
    mc_col_idx = MAIN_COLS.index("Mkt Cap ₹Cr") + 1
    for r in range(2, ws_main.max_row + 1):
        cell = ws_main.cell(r, mc_col_idx)
        if cell.value is not None:
            cell.number_format = "#,##0.00"

    # ── Sheet 2: Alias Report ──────────────────────────────────────────────
    ws_alias = wb.create_sheet("Alias Report")
    _write_sheet(ws_alias, ALIAS_COLS, alias_rows, row_fill_key="Match Method")

    # ── Sheet 3: Weight Validation ─────────────────────────────────────────
    ws_val = wb.create_sheet("Weight Validation")
    _write_sheet(ws_val, VALID_COLS, validation_rows)

    # Colour validation rows by status
    status_col_idx = VALID_COLS.index("Status") + 1
    for r in range(2, ws_val.max_row + 1):
        status = ws_val.cell(r, status_col_idx).value
        if status in STATUS_FILL:
            for cell in ws_val[r]:
                cell.fill = STATUS_FILL[status]

    # ── Sheet 4: Legend ────────────────────────────────────────────────────
    ws_leg = wb.create_sheet("Legend")
    ws_leg.append(["Colour", "Sheet", "Meaning"])
    for cell in ws_leg[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    legend_items = [
        ("Green",  "Holdings / Alias", "Exact normalised match — high confidence",  "E2EFDA"),
        ("Yellow", "Holdings / Alias", "Manual alias rule — curated by analyst",     "FFF2CC"),
        ("Blue",   "Holdings / Alias", "Fuzzy match ≥ 88 score — review advised",   "DDEBF7"),
        ("Red",    "Holdings / Alias", "No match found in AMFI data",                "FCE4D6"),
        ("Grey",   "Holdings / Alias", "Known absent from AMFI (unlisted / non-eq)", "F2F2F2"),
        ("Green",  "Validation",       "Fund weight sums within tolerance",           "E2EFDA"),
        ("Red",    "Validation",       "Fund weight below tolerance (data missing)",  "FCE4D6"),
        ("Yellow", "Validation",       "Fund weight above tolerance",                 "FFF2CC"),
    ]
    for label, sheet, meaning, hex_color in legend_items:
        ws_leg.append([label, sheet, meaning])
        ws_leg.cell(ws_leg.max_row, 1).fill = PatternFill("solid", fgColor=hex_color)

    _autowidth(ws_leg)

    # ── Sheet 5: Fund Profile ──────────────────────────────────────────────
    profile_rows, sector_cols = build_fund_profile(enriched)
    profile_cols = ["Fund"] + CAP_COLS + sector_cols + ["Sector Total"]

    ws_prof = wb.create_sheet("Fund Profile")
    ws_prof.append(profile_cols)
    for cell in ws_prof[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="top",
                                   wrap_text=True)
    ws_prof.row_dimensions[1].height = 60

    for row in profile_rows:
        ws_prof.append([row.get(c) for c in profile_cols])

    # Percentages to 2dp; blank out zero cells so the grid stays readable
    for r in range(2, ws_prof.max_row + 1):
        for c in range(2, len(profile_cols) + 1):
            cell = ws_prof.cell(r, c)
            if cell.value in (None, 0):
                cell.value = None
            else:
                cell.number_format = "0.00"

    # Freeze the fund name and header; emphasise the two total columns
    ws_prof.freeze_panes = "B2"
    for idx, name in enumerate(profile_cols, start=1):
        if name in ("Cap Total", "Sector Total"):
            for r in range(1, ws_prof.max_row + 1):
                ws_prof.cell(r, idx).font = Font(bold=True)

    _autowidth(ws_prof, max_width=14)
    ws_prof.column_dimensions["A"].width = 46

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"[writer] Output saved → {output_path}  "
          f"({len(profile_rows)} funds profiled)")
